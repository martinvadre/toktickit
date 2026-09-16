import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

REPO_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
OUT_DOCX = "/Users/martin/Uni/CPE334/Lab3/Lab Report 03.docx"
OUT_PDF = "/Users/martin/Uni/CPE334/Lab3/report_lab03_67070503466.pdf"

doc = docx.Document()

# -------------------------------------------------------------
# Page Margins & Setup (Letter 8.5 x 11 inches, 0.8 in margins)
# -------------------------------------------------------------
for sec in doc.sections:
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11.0)
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.8)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)

# -------------------------------------------------------------
# Helper Functions
# -------------------------------------------------------------
def p_title(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(20)
    run.bold = True
    return p

def p_h1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(15)
    run.bold = True
    return p

def p_h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run.bold = True
    return p

def p_body(text, bold=False, italic=False, size=10, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    return p

def p_img(rel_path, width_inches=6.3, space_after=8):
    abs_path = os.path.join(REPO_ROOT, rel_path)
    if not os.path.exists(abs_path):
        print(f"WARNING: Image not found: {abs_path}")
        return None
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(space_after)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(abs_path, width=Inches(width_inches))
    return p

def page_break():
    doc.add_page_break()

# =============================================================
# TITLE PAGE
# =============================================================
p_title("Report_lab_03")
p_body("Pawarit Wongdaeng 67070503466", bold=True, size=11)
p_body("Github : martinvadre", size=11)
p_body("Peer reviewer : Tatchakorn Thamnimitr 67070503415", bold=True, size=11)
p_body("Github : Magiciancat9", size=11)
p_body("Repository : https://github.com/martinvadre/toktickit", size=11)
p_body("Project Kanban : https://github.com/users/martinvadre/projects/3", size=11, space_after=12)

p_img("artifacts/lab-03/report/part1_github_repo_main.png", width_inches=6.3)
page_break()

# =============================================================
# ANSWER PART 1: GIT USE WITH ENGINEERING WORKFLOW
# =============================================================
p_h1("Answer Part 1: Git Use with Engineering Workflow")
p_body("Repository : https://github.com/martinvadre/toktickit")
p_body("Project : https://github.com/users/martinvadre/projects/3")
p_body("Issue 11 : https://github.com/martinvadre/toktickit/issues/29")
p_body("Issue 12 : https://github.com/martinvadre/toktickit/issues/30")
p_body("Issue 13 : https://github.com/martinvadre/toktickit/issues/31")
p_body("Issue 14 : https://github.com/martinvadre/toktickit/issues/32")
p_body("Issue 15 : https://github.com/martinvadre/toktickit/issues/33")
p_body("Issue 16 : https://github.com/martinvadre/toktickit/issues/34")
p_body("Pull request 1 : https://github.com/martinvadre/toktickit/pull/35")
p_body("Pull request 2 : https://github.com/martinvadre/toktickit/pull/36")
p_body("Pull request 3 : https://github.com/martinvadre/toktickit/pull/37")
p_body("Pull request 4 : https://github.com/martinvadre/toktickit/pull/38")
p_body("Pull request 5 : https://github.com/martinvadre/toktickit/pull/39")
p_body("Pull request 6 : https://github.com/martinvadre/toktickit/pull/40")
p_body("Pull request main : https://github.com/martinvadre/toktickit/pull/41", space_after=12)

p_h2("Kanban board")
p_img("artifacts/lab-03/report/part1_github_kanban.png", width_inches=6.3)
page_break()

p_h2("Commit History")
p_img("artifacts/lab-03/report/part1_git_commit_graph.png", width_inches=6.3)
page_break()

p_h2("All Issue")
p_img("artifacts/lab-03/report/part1_github_all_issues.png", width_inches=6.3)

p_h2("Issue 11 - https://github.com/martinvadre/toktickit/issues/29")
p_img("artifacts/lab-03/report/part1_github_issue29.png", width_inches=6.3)
page_break()

p_h2("Issue 12 - https://github.com/martinvadre/toktickit/issues/30")
p_img("artifacts/lab-03/report/part1_github_issue30.png", width_inches=6.3)

p_h2("Issue 13 - https://github.com/martinvadre/toktickit/issues/31")
p_img("artifacts/lab-03/report/part1_github_issue31.png", width_inches=6.3)
page_break()

p_h2("Issue 14 - https://github.com/martinvadre/toktickit/issues/32")
p_img("artifacts/lab-03/report/part1_github_issue32.png", width_inches=6.3)

p_h2("Issue 15 - https://github.com/martinvadre/toktickit/issues/33")
p_img("artifacts/lab-03/report/part1_github_issue33.png", width_inches=6.3)
page_break()

p_h2("Issue 16 - https://github.com/martinvadre/toktickit/issues/34")
p_img("artifacts/lab-03/report/part1_github_issue34.png", width_inches=6.3)

p_h2("All Pull Requests")
p_img("artifacts/lab-03/report/part1_github_all_prs.png", width_inches=6.3)
page_break()

p_h2("Pull Requests 1 - https://github.com/martinvadre/toktickit/pull/35")
p_img("artifacts/lab-03/report/part1_github_pr35.png", width_inches=6.3)

p_h2("Pull Requests 2 - https://github.com/martinvadre/toktickit/pull/36")
p_img("artifacts/lab-03/report/part1_github_pr36.png", width_inches=6.3)
page_break()

p_h2("Pull Requests 3 - https://github.com/martinvadre/toktickit/pull/37")
p_img("artifacts/lab-03/report/part1_github_pr37.png", width_inches=6.3)

p_h2("Pull Requests 4 - https://github.com/martinvadre/toktickit/pull/38")
p_img("artifacts/lab-03/report/part1_github_pr38.png", width_inches=6.3)
page_break()

p_h2("Pull Requests 5 - https://github.com/martinvadre/toktickit/pull/39")
p_img("artifacts/lab-03/report/part1_github_pr39.png", width_inches=6.3)

p_h2("Pull Requests 6 - https://github.com/martinvadre/toktickit/pull/40")
p_img("artifacts/lab-03/report/part1_github_pr40.png", width_inches=6.3)
page_break()

p_h2("Pull Requests main - https://github.com/martinvadre/toktickit/pull/41")
p_img("artifacts/lab-03/report/part1_github_pr41_main.png", width_inches=6.3)

p_h2("Directory Structure & .gitignore")
p_img("artifacts/lab-03/report/part1_directory_structure.png", width_inches=6.3)
p_img("artifacts/lab-03/report/part1_gitignore.png", width_inches=5.8)
page_break()

p_h2("Rendered README.md")
p_img("artifacts/lab-03/report/part1_readme_rendered.png", width_inches=6.3)
page_break()

p_h2("Rendered Reviewer.md")
p_body("Author: Pawarit Wongdaeng 67070503466 - GitHub: @martinvadre", bold=True)
p_body("Peer reviewer: Tatchakorn Thamnimitr 67070503415 - GitHub: @Magiciancat9", bold=True, space_after=8)
p_img("artifacts/lab-03/report/part1_reviewer_rendered.png", width_inches=6.3)
page_break()

# =============================================================
# ANSWER PART 2: SPEC DD
# =============================================================
p_h1("Answer Part 2: Spec DD (Specification-Driven Development)")
p_body("Specification Link: https://github.com/martinvadre/toktickit/blob/main/docs/lab-03/specification.md", bold=True)
p_body("Documented numbered Functional Requirements (FR-01 to FR-18), Business Rules (BR-01 to BR-08), Acceptance Criteria (AC-01 to AC-12), Authorization Matrix, Unified User Database Models, and Product Definition of Done before implementation began.", space_after=8)

p_h2("Rendered docs/lab-03/specification.md")
p_img("artifacts/lab-03/report/part2_spec_rendered.png", width_inches=6.3)

p_h2("Early Specification Proof: Commit 709d880 & PR #35")
p_body("PR #35 merged into lab3-staging at commit 709d880 on Sep 15, 2026, 07:33:18 prior to any feature branch implementation.")
p_img("artifacts/lab-03/report/part2_spec_early_proof.png", width_inches=6.3)
page_break()

# =============================================================
# ANSWER PART 3: TEST DD AND TRACEABILITY
# =============================================================
p_h1("Answer Part 3: Test DD and Traceability")
p_body("Test Plan Link: https://github.com/martinvadre/toktickit/blob/main/docs/lab-03/tests.md", bold=True)
p_body("Planned test matrix mapped 1:1 against Acceptance Criteria with 119 / 119 automated tests passing (100% success rate) across backend integration, frontend RTL components, and Playwright end-to-end workflows.", space_after=8)

p_h2("Rendered docs/lab-03/tests.md")
p_img("artifacts/lab-03/report/part3_tests_rendered.png", width_inches=6.3)

p_h2("Vitest Backend Suite (72 / 72 Passing)")
p_img("artifacts/lab-03/report/part3_test_backend_pass.png", width_inches=6.3)
page_break()

p_h2("Vitest Frontend Suite (40 / 40 Passing)")
p_img("artifacts/lab-03/report/part3_test_frontend_pass.png", width_inches=6.3)

p_h2("Playwright End-to-End Suite (7 / 7 Passing)")
p_img("artifacts/lab-03/report/part3_test_e2e_pass.png", width_inches=6.3)
page_break()

# =============================================================
# ANSWER PART 4: AI USE WITH REFLECTION
# =============================================================
p_h1("Answer Part 4: AI Use with Reflection")
p_body("AI Log Link: https://github.com/martinvadre/toktickit/blob/main/docs/lab-03/ai-use.md", bold=True)
p_body("Rendered view of docs/lab-03/ai-use.md documenting Gemini 3.6 Flash / Pro architecture on Google Cloud Platform, prompt iterations across 6 key development phases (PROMPT-01 to PROMPT-06), validation strategies, and engineering reflection on plan-first agent collaboration.", space_after=8)

p_img("artifacts/lab-03/report/part4_ai_use_rendered.png", width_inches=6.3)
page_break()

# =============================================================
# ANSWER PART 5: WORKING LOGIN AND PASSWORD CHANGE UI
# =============================================================
p_h1("Answer Part 5: Working Login and Password Change UI")
p_body("Demonstrates real authentication replacing the development requester selector, invalid credential rejection, inactive account lockout (AC-05), mandatory first-login password change (AC-02), and JWT session persistence.", space_after=8)

p_h2("1. Initial Login Screen")
p_img("artifacts/lab-03/screenshots/authentication/login-screen.png", width_inches=6.2)

p_h2("2. Inactive Account Lockout Banner (AC-05)")
p_img("artifacts/lab-03/screenshots/authentication/login-error-inactive.png", width_inches=6.2)
page_break()

p_h2("3. Mandatory First-Login Password Change Modal (AC-02)")
p_img("artifacts/lab-03/screenshots/authentication/password-change-mandatory.png", width_inches=6.2)

p_h2("4. Password Update Success Notification & Continuation")
p_img("artifacts/lab-03/screenshots/authentication/password-change-success.png", width_inches=6.2)
page_break()

# =============================================================
# ANSWER PART 6: WORKING IT STAFF TICKET QUEUE UI
# =============================================================
p_h1("Answer Part 6: Working IT Staff Ticket Queue UI")
p_body("Demonstrates operational ticket queue across all requesters with global keyword search, multi-criteria filtering (status, category, related system, assignee, priority), server-side sorting, pagination, and KPI counter metrics.", space_after=8)

p_h2("1. IT Staff Ticket Queue (Desktop View - AC-06)")
p_img("artifacts/lab-03/screenshots/staff-queue/queue-desktop.png", width_inches=6.3)

p_h2("2. Multi-Criteria Filter Toolbar Active State & KPI Counters")
p_img("artifacts/lab-03/screenshots/staff-queue/queue-filters.png", width_inches=6.3)
page_break()

p_h2("3. Responsive Mobile Stacked Card View (< 768px)")
p_img("artifacts/lab-03/screenshots/staff-queue/queue-mobile.png", width_inches=5.8)
page_break()

# =============================================================
# ANSWER PART 7: WORKING IT STAFF TICKET DETAIL UI
# =============================================================
p_h1("Answer Part 7: Working IT Staff Ticket Detail UI")
p_body("Demonstrates operational ticket controls (Claim/Reassign, IT Priority, Status lifecycle), mandatory resolution summary modal (AC-07), public comments vs internal staff notes with amber privacy lock badge (AC-04), and requester regression verification.", space_after=8)

p_h2("1. IT Staff Ticket Detail & Operations Overview (AC-08)")
p_img("artifacts/lab-03/screenshots/staff-ticket-detail/detail-overview.png", width_inches=6.3)

p_h2("2. Internal Staff Notes with Amber Privacy Lock Badge (AC-04)")
p_img("artifacts/lab-03/screenshots/staff-ticket-detail/detail-internal-notes.png", width_inches=6.3)
page_break()

p_h2("3. Mandatory Resolution Summary Modal Dialog (AC-07)")
p_img("artifacts/lab-03/screenshots/staff-ticket-detail/detail-resolution-modal.png", width_inches=6.3)

p_h2("4. Requester Regression View (Internal Notes Strictly Hidden)")
p_img("artifacts/lab-03/screenshots/staff-ticket-detail/requester-view-isolation.png", width_inches=6.3)
page_break()

# =============================================================
# ANSWER PART 8: WORKING ADMINISTRATOR USER MANAGEMENT UI
# =============================================================
p_h1("Answer Part 8: Working Administrator User Management UI")
p_body("Demonstrates minimalist user administration: user list with keyword search and role filtering (AC-09), create user modal (AC-11), edit user modal with self-deactivation prevented safety guardrail (AC-10), and password reset modal with mandatory first-login flag.", space_after=8)

p_h2("1. Admin User List with Search & Role Filter (AC-09)")
p_img("artifacts/lab-03/screenshots/user-management/user-list-desktop.png", width_inches=6.3)

p_h2("2. Create New User Modal (AC-11)")
p_img("artifacts/lab-03/screenshots/user-management/create-user-modal.png", width_inches=6.3)
page_break()

p_h2("3. Edit User Modal & Self-Deactivation Guardrail (AC-10)")
p_img("artifacts/lab-03/screenshots/user-management/edit-user-safety-guardrail.png", width_inches=6.3)

p_h2("4. Reset User Password Modal (AC-02 / AC-11)")
p_img("artifacts/lab-03/screenshots/user-management/reset-password-modal.png", width_inches=6.3)
page_break()

# =============================================================
# ANSWER PART 9: ZEN GREEN UI AND RESPONSIVE EVIDENCE
# =============================================================
p_h1("Answer Part 9: Zen Green UI and Responsive Evidence")
p_body("UI Spec Link: https://github.com/martinvadre/toktickit/blob/main/docs/lab-03/ui-spec.md", bold=True)
p_body("Full adherence to the Zen Green color design system (#1E5E3A, #2B8251, #EBF5EE), completed visual checklist table, and responsive behavior across Desktop (1280px), Tablet (768px), and Mobile (375px) viewports with zero horizontal overflow.", space_after=8)

p_h2("Rendered docs/lab-03/ui-spec.md")
p_img("artifacts/lab-03/report/part9_ui_spec_rendered.png", width_inches=6.3)

p_h2("Zen Green Design Verification Checklist")
p_img("artifacts/lab-03/report/part9_visual_checklist.png", width_inches=6.3)
page_break()

p_h2("Responsive Login Screen (Desktop 1280px, Tablet 768px, Mobile 375px)")
p_img("artifacts/lab-03/report/part9_responsive_login_desktop.png", width_inches=5.8)
p_img("artifacts/lab-03/report/part9_responsive_login_tablet.png", width_inches=5.0)
p_img("artifacts/lab-03/report/part9_responsive_login_mobile.png", width_inches=4.2)
page_break()

p_h2("Responsive IT Staff Queue (Desktop 1280px, Tablet 768px, Mobile 375px)")
p_img("artifacts/lab-03/report/part9_responsive_queue_desktop.png", width_inches=6.2)
p_img("artifacts/lab-03/report/part9_responsive_queue_tablet.png", width_inches=5.2)
p_img("artifacts/lab-03/report/part9_responsive_queue_mobile.png", width_inches=4.2)
page_break()

p_h2("Responsive Staff Ticket Detail (Desktop 1280px, Tablet 768px, Mobile 375px)")
p_img("artifacts/lab-03/report/part9_responsive_detail_desktop.png", width_inches=6.2)
p_img("artifacts/lab-03/report/part9_responsive_detail_tablet.png", width_inches=5.2)
p_img("artifacts/lab-03/report/part9_responsive_detail_mobile.png", width_inches=4.2)
page_break()

p_h2("Responsive Administrator User Management (Desktop 1280px, Tablet 768px, Mobile 375px)")
p_img("artifacts/lab-03/report/part9_responsive_admin_desktop.png", width_inches=6.2)
p_img("artifacts/lab-03/report/part9_responsive_admin_tablet.png", width_inches=5.2)
p_img("artifacts/lab-03/report/part9_responsive_admin_mobile.png", width_inches=4.2)

# Save DOCX
os.makedirs(os.path.dirname(OUT_DOCX), exist_ok=True)
doc.save(OUT_DOCX)
print(f"Successfully generated DOCX: {OUT_DOCX}")
