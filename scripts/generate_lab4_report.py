import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

REPO_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
OUT_DOCX = "/Users/martin/Uni/CPE334 Soft Eng/Lab4/Lab Report 04.docx"
OUT_PDF = "/Users/martin/Uni/CPE334 Soft Eng/Lab4/Lab Report 04.pdf"

doc = docx.Document()

# -------------------------------------------------------------
# Page Margins & Setup (Letter 8.5 x 11 inches, 0.8 in margins)
# -------------------------------------------------------------
for sec in doc.sections:
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11.0)
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.8)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)

# -------------------------------------------------------------
# Helper Functions
# -------------------------------------------------------------
def p_title(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(20)
    run.bold = True
    return p

def p_h1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(15)
    run.bold = True
    return p

def p_h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run.bold = True
    return p

def p_body(text, bold=False, italic=False, size=10, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    return p

def p_img(rel_path, width_inches=5.6, height_inches=None, space_after=6):
    abs_path = os.path.join(REPO_ROOT, rel_path)
    if not os.path.exists(abs_path):
        print(f"WARNING: Image not found: {abs_path}")
        return None
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(space_after)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    if height_inches is not None:
        run.add_picture(abs_path, height=Inches(height_inches))
    else:
        run.add_picture(abs_path, width=Inches(width_inches))
    return p

def page_break():
    doc.add_page_break()

# =============================================================
# TITLE PAGE
# =============================================================
p_title("Report_lab_04")
p_body("Pawarit Wongdaeng 67070503466", bold=True, size=11)
p_body("Github : martinvadre", size=11)
p_body("Peer reviewer : Tatchakorn Thamnimitr 67070503415", bold=True, size=11)
p_body("Github : Magiciancat9", size=11)
p_body("Repository : https://github.com/martinvadre/toktickit", size=11)
p_body("Project Kanban : https://github.com/users/martinvadre/projects/3", size=11, space_after=12)

p_img("artifacts/lab-04/report/part1_github_repo_main.png", width_inches=6.3)
page_break()

# =============================================================
# ANSWER PART 1: GIT USE WITH ENGINEERING WORKFLOW
# =============================================================
p_h1("Answer Part 1: Git Use with Engineering Workflow")
p_body("Repository : https://github.com/martinvadre/toktickit")
p_body("Project : https://github.com/users/martinvadre/projects/3")
p_body("Issue 17 : https://github.com/martinvadre/toktickit/issues/55")
p_body("Issue 18 : https://github.com/martinvadre/toktickit/issues/56")
p_body("Issue 19 : https://github.com/martinvadre/toktickit/issues/57")
p_body("Issue 20 : https://github.com/martinvadre/toktickit/issues/58")
p_body("Issue 21 : https://github.com/martinvadre/toktickit/issues/68")
p_body("Issue 22 : https://github.com/martinvadre/toktickit/issues/71")
p_body("Pull request 1 : https://github.com/martinvadre/toktickit/pull/62")
p_body("Pull request 2 : https://github.com/martinvadre/toktickit/pull/63")
p_body("Pull request 3 : https://github.com/martinvadre/toktickit/pull/64")
p_body("Pull request 4 : https://github.com/martinvadre/toktickit/pull/65")
p_body("Pull request 5 : https://github.com/martinvadre/toktickit/pull/70")
p_body("Pull request 6 : https://github.com/martinvadre/toktickit/pull/72")
p_body("Pull request main : https://github.com/martinvadre/toktickit/pull/73")
p_body("Peer review PR 1 : https://github.com/Magiciancat9/toktickit/pull/60")
p_body("Peer review PR 2 : https://github.com/Magiciancat9/toktickit/pull/63")
p_body("Peer review PR 3 : https://github.com/Magiciancat9/toktickit/pull/65")
p_body("Peer review PR 4 : https://github.com/Magiciancat9/toktickit/pull/67")
p_body("Peer review PR 5 : https://github.com/Magiciancat9/toktickit/pull/76")
p_body("Peer review PR 6 : https://github.com/Magiciancat9/toktickit/pull/77")
p_body("Peer review PR 7 : https://github.com/Magiciancat9/toktickit/pull/78")
p_body("Peer review PR 8 : https://github.com/Magiciancat9/toktickit/pull/79")
p_body("Peer review PR 9 : https://github.com/Magiciancat9/toktickit/pull/80")
p_body("Peer review PR 10 : https://github.com/Magiciancat9/toktickit/pull/81")
p_body("Peer review PR 11 : https://github.com/Magiciancat9/toktickit/pull/82")
p_body("Peer review PR 12 : https://github.com/Magiciancat9/toktickit/pull/83")
p_body("Peer review PR main : https://github.com/Magiciancat9/toktickit/pull/84", space_after=12)
page_break()

p_h2("Kanban board")
p_img("artifacts/lab-04/report/part1_github_kanban.png", width_inches=5.5)
page_break()

p_h2("Commit History")
p_img("artifacts/lab-04/report/part1_git_commit_graph.png", width_inches=6.2)
page_break()

p_h2("All Issue")
p_img("artifacts/lab-04/report/part1_github_all_issues.png", width_inches=5.5)

p_h2("Issue 17 - https://github.com/martinvadre/toktickit/issues/55")
p_img("artifacts/lab-04/report/part1_github_issue55.png", width_inches=5.5)
page_break()

p_h2("Issue 18 - https://github.com/martinvadre/toktickit/issues/56")
p_img("artifacts/lab-04/report/part1_github_issue56.png", width_inches=5.5)

p_h2("Issue 19 - https://github.com/martinvadre/toktickit/issues/57")
p_img("artifacts/lab-04/report/part1_github_issue57.png", width_inches=5.5)
page_break()

p_h2("Issue 20 - https://github.com/martinvadre/toktickit/issues/58")
p_img("artifacts/lab-04/report/part1_github_issue58.png", width_inches=5.5)

p_h2("Issue 21 - https://github.com/martinvadre/toktickit/issues/68")
p_img("artifacts/lab-04/report/part1_github_issue68.png", width_inches=5.5)
page_break()

p_h2("Issue 22 - https://github.com/martinvadre/toktickit/issues/71")
p_img("artifacts/lab-04/report/part1_github_issue71.png", width_inches=5.5)

p_h2("All Pull Requests")
p_img("artifacts/lab-04/report/part1_github_all_prs.png", width_inches=5.5)
page_break()

p_h2("Pull Requests 1 - https://github.com/martinvadre/toktickit/pull/62")
p_img("artifacts/lab-04/report/part1_github_pr62.png", width_inches=5.5)

p_h2("Pull Requests 2 - https://github.com/martinvadre/toktickit/pull/63")
p_img("artifacts/lab-04/report/part1_github_pr63.png", width_inches=5.5)
page_break()

p_h2("Pull Requests 3 - https://github.com/martinvadre/toktickit/pull/64")
p_img("artifacts/lab-04/report/part1_github_pr64.png", width_inches=5.5)

p_h2("Pull Requests 4 - https://github.com/martinvadre/toktickit/pull/65")
p_img("artifacts/lab-04/report/part1_github_pr65.png", width_inches=5.5)
page_break()

p_h2("Pull Requests 5 - https://github.com/martinvadre/toktickit/pull/70")
p_img("artifacts/lab-04/report/part1_github_pr70.png", width_inches=5.5)

p_h2("Pull Requests 6 - https://github.com/martinvadre/toktickit/pull/72")
p_img("artifacts/lab-04/report/part1_github_pr72.png", width_inches=5.5)
page_break()

p_h2("Pull Requests main - https://github.com/martinvadre/toktickit/pull/73")
p_img("artifacts/lab-04/report/part1_github_pr73.png", width_inches=5.5)

p_h2("Peer Review PR 1 (Partner @Magiciancat9) - https://github.com/Magiciancat9/toktickit/pull/60")
p_img("artifacts/lab-04/report/part1_partner_pr60.png", width_inches=5.5)
page_break()

p_h2("Peer Review PR 2 (Partner @Magiciancat9) - https://github.com/Magiciancat9/toktickit/pull/63")
p_img("artifacts/lab-04/report/part1_partner_pr63.png", width_inches=5.5)

p_h2("Peer Review PR 3 (Partner @Magiciancat9) - https://github.com/Magiciancat9/toktickit/pull/65")
p_img("artifacts/lab-04/report/part1_partner_pr65.png", width_inches=5.5)
page_break()

p_h2("Peer Review PR 4 (Partner @Magiciancat9) - https://github.com/Magiciancat9/toktickit/pull/67")
p_img("artifacts/lab-04/report/part1_partner_pr67.png", width_inches=5.5)

p_h2("Peer Review PR 5 (Partner @Magiciancat9) - https://github.com/Magiciancat9/toktickit/pull/76")
p_img("artifacts/lab-04/report/part1_partner_pr76.png", width_inches=5.5)
page_break()

p_h2("Peer Review PR 6 (Partner @Magiciancat9) - https://github.com/Magiciancat9/toktickit/pull/77")
p_img("artifacts/lab-04/report/part1_partner_pr77.png", width_inches=5.5)

p_h2("Peer Review PR 7 (Partner @Magiciancat9) - https://github.com/Magiciancat9/toktickit/pull/78")
p_img("artifacts/lab-04/report/part1_partner_pr78.png", width_inches=5.5)
page_break()

p_h2("Peer Review PR 8 (Partner @Magiciancat9) - https://github.com/Magiciancat9/toktickit/pull/79")
p_img("artifacts/lab-04/report/part1_partner_pr79.png", width_inches=5.5)

p_h2("Peer Review PR 9 (Partner @Magiciancat9) - https://github.com/Magiciancat9/toktickit/pull/80")
p_img("artifacts/lab-04/report/part1_partner_pr80.png", width_inches=5.5)
page_break()

p_h2("Peer Review PR 10 (Partner @Magiciancat9) - https://github.com/Magiciancat9/toktickit/pull/81")
p_img("artifacts/lab-04/report/part1_partner_pr81.png", width_inches=5.5)

p_h2("Peer Review PR 11 (Partner @Magiciancat9) - https://github.com/Magiciancat9/toktickit/pull/82")
p_img("artifacts/lab-04/report/part1_partner_pr82.png", width_inches=5.5)
page_break()

p_h2("Peer Review PR 12 (Partner @Magiciancat9) - https://github.com/Magiciancat9/toktickit/pull/83")
p_img("artifacts/lab-04/report/part1_partner_pr83.png", width_inches=5.5)

p_h2("Peer Review PR main (Partner @Magiciancat9) - https://github.com/Magiciancat9/toktickit/pull/84")
p_img("artifacts/lab-04/report/part1_partner_pr84.png", width_inches=5.5)
page_break()

p_h2("Directory Structure & .gitignore")
p_img("artifacts/lab-04/report/part1_directory_structure.png", width_inches=5.8)
p_img("artifacts/lab-04/report/part1_gitignore.png", width_inches=5.5)
page_break()

p_h2("Rendered README.md")
p_img("artifacts/lab-04/report/part1_readme_rendered.png", width_inches=6.2)
page_break()

p_h2("Rendered Reviewer.md")
p_body("Author: Pawarit Wongdaeng 67070503466 - GitHub: @martinvadre", bold=True)
p_body("Peer reviewer: Tatchakorn Thamnimitr 67070503415 - GitHub: @Magiciancat9", bold=True, space_after=6)
p_img("artifacts/lab-04/report/part1_reviewer_rendered.png", width_inches=5.4, space_after=4)
p_img("artifacts/lab-04/report/part1_reviewer_rendered_2.png", width_inches=5.4, space_after=4)
page_break()

p_h2("Rendered Reviewer.md (Partner PR Reviews Continued)")
p_img("artifacts/lab-04/report/part1_reviewer_rendered_3.png", width_inches=6.2)
page_break()

# =============================================================
# ANSWER PART 2: SPEC DD
# =============================================================
p_h1("Answer Part 2: Spec DD (Specification-Driven Development)")
p_body("Specification Link: https://github.com/martinvadre/toktickit/blob/main/docs/lab-04/specification.md", bold=True, space_after=6)

p_h2("Rendered docs/lab-04/specification.md")
p_img("artifacts/lab-04/report/part2_spec_rendered.png", height_inches=3.0)

p_h2("Early Specification Proof: Commit 6eae025 & PR #62")
p_img("artifacts/lab-04/report/part2_spec_early_proof.png", height_inches=3.0)
page_break()

# =============================================================
# ANSWER PART 3: TEST DD AND TRACEABILITY
# =============================================================
p_h1("Answer Part 3: Test DD and Traceability")
p_body("Test Plan Link: https://github.com/martinvadre/toktickit/blob/main/docs/lab-04/tests.md", bold=True, space_after=6)

p_h2("Rendered docs/lab-04/tests.md")
p_img("artifacts/lab-04/report/part3_tests_rendered.png", height_inches=3.0)

p_h2("Vitest Backend Suite (92 / 92 Passing)")
p_img("artifacts/lab-04/report/part3_test_backend_pass.png", height_inches=3.0)
page_break()

p_h2("Vitest Frontend Suite (53 / 53 Passing)")
p_img("artifacts/lab-04/report/part3_test_frontend_pass.png", width_inches=5.5)

p_h2("Playwright End-to-End Suite (14 / 14 Passing)")
p_img("artifacts/lab-04/report/part3_test_e2e_pass.png", width_inches=5.5)
page_break()

# =============================================================
# ANSWER PART 4: AI USE WITH REFLECTION
# =============================================================
p_h1("Answer Part 4: AI Use with Reflection")
p_body("AI Log Link: https://github.com/martinvadre/toktickit/blob/main/docs/lab-04/ai-use.md", bold=True, space_after=8)

p_img("artifacts/lab-04/report/part4_ai_use_rendered.png", width_inches=6.2)
page_break()

# =============================================================
# ANSWER PART 5: WORKING IT STAFF DASHBOARD UI
# =============================================================
p_h1("Answer Part 5: Working IT Staff Dashboard UI")

p_h2("1. IT Staff Dashboard (Desktop View 1280px)")
p_img("artifacts/lab-04/screenshots/staff-dashboard/desktop.png", height_inches=3.0)

p_h2("2. Database Query Proof")
p_img("artifacts/lab-04/report/part5_staff_dashboard_sql_proof.png", height_inches=3.0)
page_break()

p_h2("3. IT Staff Dashboard (Tablet View 768px)")
p_img("artifacts/lab-04/screenshots/staff-dashboard/tablet.png", height_inches=3.4)

p_h2("4. IT Staff Dashboard (Mobile View 375px)")
p_img("artifacts/lab-04/screenshots/staff-dashboard/mobile.png", height_inches=3.4)
page_break()

# =============================================================
# ANSWER PART 6: WORKING ACTIONS TAKEN UI
# =============================================================
p_h1("Answer Part 6: Working Actions Taken UI")

p_h2("1. IT Staff Ticket Detail with Actions Taken List (Desktop)")
p_img("artifacts/lab-04/screenshots/actions-taken/staff-actions-desktop.png", height_inches=3.0)

p_h2("2. Record / Edit Action Taken Modal")
p_img("artifacts/lab-04/screenshots/actions-taken/action-modal.png", height_inches=3.0)
page_break()

p_h2("3. Requester Ticket Detail (Read-Only Actions Taken View)")
p_img("artifacts/lab-04/screenshots/actions-taken/requester-actions-desktop.png", height_inches=5.6)
page_break()

# =============================================================
# ANSWER PART 7: WORKING TICKET WORKFLOW
# =============================================================
p_h1("Answer Part 7: Working Ticket Workflow")

p_h2("1. Mandatory Resolution Summary Modal")
p_img("artifacts/lab-04/report/part7_resolution_modal.png", height_inches=3.0)

p_h2("2. Requester Resolution Indication Badge")
p_img("artifacts/lab-04/report/part7_advisory_cue.png", height_inches=3.0)
page_break()

# =============================================================
# ANSWER PART 8: WORKING REQUESTER DASHBOARD AND FINAL REGRESSION UI
# =============================================================
p_h1("Answer Part 8: Working Requester Dashboard and Final Regression UI")

p_h2("1. Requester Dashboard (Desktop View 1280px)")
p_img("artifacts/lab-04/screenshots/requester-dashboard/desktop.png", height_inches=3.0)

p_h2("2. Requester Dashboard (Tablet View 768px)")
p_img("artifacts/lab-04/screenshots/requester-dashboard/tablet.png", height_inches=3.0)
page_break()

p_h2("3. Requester Dashboard (Mobile View 375px)")
p_img("artifacts/lab-04/screenshots/requester-dashboard/mobile.png", height_inches=5.6)
page_break()

# =============================================================
# ANSWER PART 9: ZEN GREEN UI, RESPONSIVE, ACCESSIBILITY, AND FINAL POLISH
# =============================================================
p_h1("Answer Part 9: Zen Green UI, Responsive, Accessibility, and Final Polish")
p_body("UI Spec Link: https://github.com/martinvadre/toktickit/blob/main/docs/lab-04/ui-spec.md", bold=True, space_after=6)

p_h2("Rendered docs/lab-04/ui-spec.md")
p_img("artifacts/lab-04/report/part9_ui_spec_rendered.png", height_inches=3.0)

p_h2("Zen Green Design & Accessibility Verification Checklist")
p_img("artifacts/lab-04/report/part9_visual_checklist.png", height_inches=3.0)

# Save DOCX
os.makedirs(os.path.dirname(OUT_DOCX), exist_ok=True)
doc.save(OUT_DOCX)
print(f"Successfully generated DOCX: {OUT_DOCX}")
print("Note: PDF is compiled using scripts/generate_lab4_pdf.js via Playwright for pixel-perfect fidelity.")
